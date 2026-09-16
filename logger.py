from pathlib import Path
from openpyxl import Workbook, load_workbook
from datetime import datetime
HEADERS=['Date','Time','Tweet ID','Author','Tweet','Tweet URL','Tweet Age','Likes','Replies','Search Keyword','Generated Reply','Reply Status','Reply URL','Error']
class ExcelLogger:
    def __init__(self,path:Path):
        self.path=path; path.parent.mkdir(parents=True,exist_ok=True)
        if not path.exists():
            wb=Workbook(); ws=wb.active; ws.title='Replies'; ws.append(HEADERS); wb.save(path)
    def log(self, t, keyword, reply='', status='', reply_url='', error='', age=''):
        wb=load_workbook(self.path); ws=wb['Replies']; now=datetime.now(); ws.append([now.date().isoformat(),now.strftime('%H:%M:%S'),str(t.id),t.author,t.text,t.url,age,t.likes,t.replies,keyword,reply,status,reply_url,error]); wb.save(self.path)
